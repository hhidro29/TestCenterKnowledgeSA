from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


OUTPUT = "/Users/fa-2400/Documents/ChatGPT/Knowledge Management/Ringkasan Masukan Kak Nur.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_font(run, name="Aptos", size=10.5, bold=False, color="000000", italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_font(r)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_font(r)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, size=14 if level == 1 else 11.5, bold=True, color="000000")
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_shading(cell, "1F4E79")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_font(r, size=9.5, bold=True, color="FFFFFF")
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2 == 1:
                set_cell_shading(cell, "F2F6FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(val)
            set_font(r, size=9.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for style_name, size in (("Title", 23), ("Heading 1", 14), ("Heading 2", 11.5)):
        st = styles[style_name]
        st.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        st._element.rPr.rFonts.set(qn("w:ascii"), st.font.name)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), st.font.name)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        p_pr = st._element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)
    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(10.5)
    styles["List Bullet 2"].font.name = "Aptos"
    styles["List Bullet 2"].font.size = Pt(10.5)
    styles["List Number"].font.name = "Aptos"
    styles["List Number"].font.size = Pt(10.5)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("Ringkasan Masukan Kak Nur")
    set_font(r, size=8.5, color="666666")


def build():
    doc = Document()
    configure_styles(doc)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    add_footer(sec)

    # Title block
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(6)
    remove_paragraph_borders(p)
    r = p.add_run("Ringkasan Masukan Kak Nur")
    set_font(r, name="Aptos Display", size=23, bold=True, color="000000")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(18)
    remove_paragraph_borders(p2)
    r = p2.add_run("Pemantauan siswa, konsultasi orang tua, data pembelajaran, dan kebutuhan tindak lanjut")
    set_font(r, size=11.5, color="000000")

    add_body(doc, "Dokumen ini merangkum pengalaman dan kebutuhan Kak Nur dalam mendampingi siswa serta berkomunikasi dengan orang tua pada jenjang SD–SMP dan SMA. Temuan utamanya adalah perlunya pemantauan siswa yang lebih cepat dan terintegrasi, informasi akademik yang mudah diakses, materi yang lebih lengkap, serta panduan penanganan kasus dan keluhan yang lebih jelas.")

    add_heading(doc, "Gambaran Umum", 1)
    add_body(doc, "Interaksi dengan orang tua berlangsung secara situasional. Pada jenjang SMA, orang tua paling sering menanyakan progres nilai, peluang masuk jurusan atau universitas tertentu, dan langkah yang perlu dilakukan agar nilai siswa meningkat. Pada jenjang SD–SMP, pertanyaan lebih banyak berkaitan dengan kehadiran, jadwal, penggunaan aplikasi, progres per mata pelajaran, dan persiapan masuk sekolah tertentu.")
    add_body(doc, "Kak Nur memantau keaktifan siswa melalui kehadiran harian, Tryout di BA, drill soal, kegiatan les, klinik PR, dan penugasan. Karena laporan dari pusat terkadang cukup lama, siswa diminta mengirimkan screenshot hasil latihan atau Tryout jika data dibutuhkan segera.")

    add_heading(doc, "Temuan Utama", 1)
    add_heading(doc, "Aktivitas dan kehadiran siswa", 2)
    add_bullet(doc, "Kegiatan yang dipantau meliputi Tryout di BA, drill soal, les, klinik PR, penugasan, dan kehadiran harian.")
    add_bullet(doc, "Pada jenjang bawah, kebutuhan orang tua lebih banyak terkait teknis, seperti jadwal, kehadiran, login, dan penggunaan aplikasi.")
    add_bullet(doc, "Jika siswa tidak dapat mengikuti les, jadwal dapat dialihkan ke klinik PR.")
    add_bullet(doc, "Siswa jenjang atas relatif jarang meminta perubahan jadwal.")
    add_bullet(doc, "Kak Nur berharap siswa lebih aktif karena sebagian siswa masih jarang hadir atau belum konsisten mengikuti kegiatan.")
    add_bullet(doc, "Target siswa lebih diarahkan pada keaktifan dan konsistensi harian, dengan contoh target latihan 60 soal per minggu.")

    add_heading(doc, "Komunikasi dengan orang tua dan siswa", 2)
    add_bullet(doc, "Orang tua SMA kelas 10–11 biasanya menghubungi secara acak untuk menanyakan progres belajar dan target pendidikan.")
    add_bullet(doc, "Pertanyaan yang sering muncul adalah target nilai untuk jurusan atau universitas tertentu, peluang masuk kampus, program peningkatan nilai, dan status kelulusan.")
    add_bullet(doc, "Pertanyaan teknis dapat dijelaskan secara garis besar melalui WhatsApp, sedangkan pembahasan target kuliah diarahkan ke sesi konsultasi yang dijadwalkan oleh SA.")
    add_bullet(doc, "Pertanyaan dari siswa diprioritaskan karena berkaitan langsung dengan kebutuhan belajar dan tindak lanjut harian.")
    add_bullet(doc, "Orang tua SD–SMP perlu mendapat laporan rutin mengenai kehadiran, aktivitas belajar, dan progres per mata pelajaran.")

    add_heading(doc, "Pemantauan nilai dan progres siswa", 2)
    add_body(doc, "Dalam konsultasi, Kak Nur biasanya memulai dari nilai per subtes, kemudian membandingkannya dengan hasil Tryout, nilai sekolah, data profiling, konsistensi mengikuti les, dan data penerimaan tahun sebelumnya. Jika siswa sudah memiliki nilai tinggi tetapi tidak mengalami kenaikan, kondisi tersebut disampaikan secara terbuka dan dibahas bersama untuk menentukan nilai yang masih perlu dikejar.")
    add_bullet(doc, "Mencari data siswa tahun sebelumnya yang diterima di jurusan atau universitas yang dituju.")
    add_bullet(doc, "Membandingkan data yang dimiliki Ruangguru dengan data Zebracross.")
    add_bullet(doc, "Menggunakan data nilai minimum atau nilai penerimaan tahun sebelumnya sebagai gambaran target.")
    add_bullet(doc, "Merekomendasikan jurusan atau universitas alternatif jika target awal terlalu sulit dicapai.")
    add_bullet(doc, "Meminta screenshot hasil drill atau Tryout ketika laporan resmi dari pusat belum tersedia dan informasinya dibutuhkan secara urgent.")

    add_heading(doc, "Kebutuhan siswa SMA", 2)
    add_bullet(doc, "Orang tua membutuhkan jawaban yang konkret mengenai posisi nilai siswa terhadap target jurusan atau universitas.")
    add_bullet(doc, "Setelah Tryout, siswa perlu mendapatkan target lanjutan yang jelas, bukan hanya hasil nilai.")
    add_bullet(doc, "Siswa yang belum mencapai target perlu diarahkan ke klinik topik, klinik PR, drill soal tambahan, atau penugasan berdasarkan subtes yang masih rendah.")
    add_bullet(doc, "Diperlukan rekomendasi program yang lebih terstruktur berdasarkan kekurangan nilai setiap siswa.")

    add_heading(doc, "Kebutuhan siswa SD dan SMP", 2)
    add_bullet(doc, "Orang tua menanyakan progres pembelajaran per mata pelajaran, kehadiran, jadwal, dan cara login akun siswa.")
    add_bullet(doc, "Orang tua juga menanyakan persiapan masuk sekolah tertentu, misalnya jalur MTs ke MAN 2 atau sekolah favorit lainnya.")
    add_bullet(doc, "Saat ini belum tersedia program khusus untuk membantu siswa masuk ke sekolah tertentu.")
    add_bullet(doc, "Dibutuhkan panduan pilihan sekolah dan program persiapan yang lebih jelas untuk jenjang bawah.")

    add_heading(doc, "Data dan dashboard", 1)
    add_body(doc, "Data Studio sebelumnya digunakan untuk melihat kehadiran, nilai kuis dan Tryout, materi yang telah dipelajari, perkembangan nilai, serta perbandingan nilai melalui grafik. Ketika data atau laporan belum tersedia, Kak Nur perlu melakukan tracking manual dan membuat report card sendiri.")
    add_body(doc, "Dashboard yang dibutuhkan perlu membantu Kak Nur menjawab pertanyaan orang tua dengan cepat, bukan hanya menampilkan data mentah. Data sebaiknya dapat dilihat per siswa, per subtes, dan dari waktu ke waktu.")
    add_table(doc, ["Kebutuhan dashboard", "Manfaat dalam konsultasi"], [
        ("Perkembangan nilai dalam bentuk grafik", "Menunjukkan tren kenaikan atau stagnasi nilai secara mudah."),
        ("Perbandingan nilai per subtes", "Menentukan area belajar yang paling membutuhkan intervensi."),
        ("Target jurusan atau universitas", "Membandingkan posisi nilai siswa dengan target yang relevan."),
        ("Kehadiran dan aktivitas harian", "Menilai konsistensi mengikuti les, drill, Tryout, dan penugasan."),
        ("Riwayat Tryout dan data tahun sebelumnya", "Menjadi dasar diskusi target dan rekomendasi alternatif."),
        ("Progress report yang lebih cepat", "Mengurangi ketergantungan pada screenshot dan tracking manual."),
    ], widths=[2.25, 4.35])

    add_heading(doc, "Pertanyaan yang sulit dijawab", 1)
    add_body(doc, "Pertanyaan umum yang informasinya tersedia masih dapat ditangani. Tantangan muncul ketika orang tua menanyakan kuota, persyaratan, jalur kedinasan, atau informasi institusi yang belum dirilis. Untuk menjawabnya, Kak Nur mencari informasi melalui Linktree, internet, dan website resmi institusi, tetapi terdapat risiko informasi belum terbaru.")
    add_bullet(doc, "Diperlukan sumber informasi resmi dan terpusat mengenai PTN, jalur mandiri, SNBP, sekolah kedinasan, kuota, dan persyaratan.")
    add_bullet(doc, "Jika informasi belum tersedia, perlu ada alur eskalasi atau PIC yang dapat membantu pencarian dan validasi data.")
    add_bullet(doc, "Jawaban terkait peluang masuk universitas sebaiknya tetap dikaitkan dengan progres dan nilai siswa, bukan hanya informasi umum.")

    add_heading(doc, "Kendala materi pembelajaran", 1)
    add_bullet(doc, "Tema atau urutan bab di Ruangbelajar terkadang berbeda dengan buku sekolah siswa, sehingga perlu dicari padanan nama bab terlebih dahulu.")
    add_bullet(doc, "Orang tua mengeluhkan materi yang lebih banyak berbentuk video dan meminta materi dalam bentuk PDF yang dapat diunduh atau dicetak.")
    add_bullet(doc, "Materi dinilai belum selalu lengkap, terutama untuk siswa IPS.")
    add_bullet(doc, "Permintaan materi cetak paling banyak berasal dari siswa SMA, khususnya kelas 11–12 dan jurusan IPS.")
    add_bullet(doc, "Rangkuman materi diharapkan mencakup keseluruhan topik, bukan hanya sebagian materi.")

    add_heading(doc, "Kasus layanan dan penanganan keluhan", 1)
    add_body(doc, "Beberapa keluhan memerlukan koordinasi lintas tim dan respons yang cepat. Kak Nur pernah menghadapi siswa yang khawatir tidak lolos SNBP atau jalur mandiri, siswa yang terkendala cicilan, serta laporan teknis yang belum menghasilkan perbaikan. Pada kasus tertentu, Kak Nur perlu meminta arahan kepada rekan atau atasan karena belum tersedia panduan penanganan yang seragam.")
    add_table(doc, ["Kasus atau keluhan", "Kebutuhan tindak lanjut"], [
        ("Siswa khawatir tidak lolos SNBP atau jalur mandiri", "Panduan komunikasi, alternatif jalur atau kampus, dan alur konsultasi."),
        ("Siswa terkendala cicilan dan tidak dapat melanjutkan layanan", "Kejelasan status akses, kebijakan bantuan, dan PIC penyelesaian kasus."),
        ("Laporan sudah disampaikan tetapi belum ada perbaikan", "Status tiket, batas waktu respons, dan kepastian eskalasi."),
        ("Aplikasi sering bug atau harus di-refresh", "Validasi masalah, pembaruan status, dan informasi solusi sementara."),
        ("Respons pusat untuk kasus urgent cukup lama", "Kanal prioritas dan SLA untuk kasus yang berdampak langsung pada siswa."),
    ], widths=[2.65, 3.95])

    add_heading(doc, "Kebutuhan prioritas", 1)
    add_body(doc, "Berdasarkan seluruh masukan, berikut kebutuhan yang paling berdampak terhadap pekerjaan pendampingan siswa dan kualitas komunikasi dengan orang tua:")
    priorities = [
        "Menyediakan dashboard terintegrasi untuk kehadiran, drill, Tryout, nilai, materi, dan grafik perkembangan siswa.",
        "Mempercepat ketersediaan hasil Tryout dan progress report, terutama untuk kebutuhan konsultasi yang urgent.",
        "Menyediakan rekomendasi program belajar berdasarkan subtes atau materi yang masih lemah.",
        "Membangun database nilai penerimaan dan kelulusan tahun sebelumnya yang mudah dicari berdasarkan jurusan dan universitas.",
        "Menyediakan sumber informasi resmi mengenai PTN, jalur mandiri, SNBP, kedinasan, kuota, dan persyaratan sekolah.",
        "Membuat panduan penanganan complaint dan eskalasi kasus, termasuk kendala cicilan, akses layanan, dan masalah teknis.",
        "Menyediakan materi PDF yang lengkap dan dapat diunduh atau dicetak, terutama untuk SMA kelas 11–12 dan jurusan IPS.",
        "Mengembangkan program persiapan masuk sekolah tertentu bagi siswa SD–SMP.",
        "Menetapkan pola komunikasi rutin kepada orang tua mengenai kehadiran, aktivitas, nilai, dan perkembangan siswa.",
    ]
    for item in priorities:
        add_numbered(doc, item)

    add_heading(doc, "Kesimpulan", 1)
    add_body(doc, "Peran pendamping saat ini sangat bergantung pada tracking manual, pencarian data lintas sumber, dan koordinasi personal ketika muncul kasus khusus. Dengan dashboard yang lebih lengkap, sumber informasi yang terpusat, materi PDF yang memadai, serta alur eskalasi yang jelas, pendamping dapat memberikan jawaban yang lebih cepat, konsisten, dan berbasis data kepada siswa maupun orang tua.")

    # Set core properties
    props = doc.core_properties
    props.title = "Ringkasan Masukan Kak Nur"
    props.subject = "Ringkasan kebutuhan pendampingan siswa dan komunikasi orang tua"
    props.author = ""
    props.comments = ""
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
